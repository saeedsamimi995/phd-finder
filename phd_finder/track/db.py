from __future__ import annotations

import json
import sqlite3
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font

from phd_finder.models import Position
from phd_finder.paths import DB_PATH, OUTPUT_DIR, ensure_dirs

COLUMNS = [
    "id",
    "title",
    "institution",
    "city",
    "country",
    "region_tier",
    "pi_or_contact",
    "email",
    "url",
    "source",
    "deadline",
    "posted_date",
    "application_mode",
    "match_reasons",
    "score",
    "priority_rank",
    "status",
    "package_path",
    "notes",
    "topic_score",
    "region_weight",
    "discovered_at",
    "description",
]


class Database:
    def __init__(self, path: Path | None = None) -> None:
        ensure_dirs()
        self.path = path or DB_PATH
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._init_schema()

    def connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.path)
        conn.row_factory = sqlite3.Row
        return conn

    def _init_schema(self) -> None:
        with self.connect() as conn:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS positions (
                    id TEXT PRIMARY KEY,
                    title TEXT NOT NULL,
                    institution TEXT,
                    city TEXT,
                    country TEXT,
                    region_tier TEXT,
                    pi_or_contact TEXT,
                    email TEXT,
                    url TEXT,
                    source TEXT,
                    deadline TEXT,
                    posted_date TEXT,
                    application_mode TEXT,
                    match_reasons TEXT,
                    score REAL,
                    priority_rank INTEGER,
                    status TEXT,
                    package_path TEXT,
                    notes TEXT,
                    topic_score REAL,
                    region_weight REAL,
                    discovered_at TEXT,
                    description TEXT,
                    raw_json TEXT
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS applied_history (
                    key TEXT PRIMARY KEY,
                    label TEXT,
                    url TEXT,
                    imported_at TEXT
                )
                """
            )
            conn.commit()

    def upsert_positions(self, positions: Iterable[Position]) -> int:
        count = 0
        with self.connect() as conn:
            for pos in positions:
                existing = conn.execute(
                    "SELECT status, package_path, notes, priority_rank FROM positions WHERE id = ?",
                    (pos.id,),
                ).fetchone()
                if existing:
                    # Keep application progress. Matching status (new/skipped/closed) can be recomputed.
                    keep = {"docs_ready", "awaiting_approval", "approved", "email_sent", "portal_ready", "applied"}
                    status = existing["status"] if existing["status"] in keep else pos.status
                    package_path = existing["package_path"] or pos.package_path
                    notes = existing["notes"] or pos.notes
                    priority_rank = pos.priority_rank
                else:
                    status = pos.status
                    package_path = pos.package_path
                    notes = pos.notes
                    priority_rank = pos.priority_rank

                conn.execute(
                    """
                    INSERT INTO positions (
                        id, title, institution, city, country, region_tier, pi_or_contact,
                        email, url, source, deadline, posted_date, application_mode,
                        match_reasons, score, priority_rank, status, package_path, notes,
                        topic_score, region_weight, discovered_at, description, raw_json
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(id) DO UPDATE SET
                        title=excluded.title,
                        institution=excluded.institution,
                        city=excluded.city,
                        country=excluded.country,
                        region_tier=excluded.region_tier,
                        pi_or_contact=excluded.pi_or_contact,
                        email=excluded.email,
                        url=excluded.url,
                        source=excluded.source,
                        deadline=excluded.deadline,
                        posted_date=excluded.posted_date,
                        application_mode=excluded.application_mode,
                        match_reasons=excluded.match_reasons,
                        score=excluded.score,
                        priority_rank=?,
                        status=?,
                        package_path=?,
                        notes=?,
                        topic_score=excluded.topic_score,
                        region_weight=excluded.region_weight,
                        description=excluded.description,
                        raw_json=excluded.raw_json
                    """,
                    (
                        pos.id,
                        pos.title,
                        pos.institution,
                        pos.city,
                        pos.country,
                        pos.region_tier,
                        pos.pi_or_contact,
                        pos.email,
                        pos.url,
                        pos.source,
                        pos.deadline,
                        pos.posted_date,
                        pos.application_mode,
                        pos.match_reasons,
                        pos.score,
                        priority_rank,
                        status,
                        package_path,
                        notes,
                        pos.topic_score,
                        pos.region_weight,
                        pos.discovered_at,
                        pos.description[:20000] if pos.description else "",
                        json.dumps(pos.raw, ensure_ascii=False)[:50000],
                        priority_rank,
                        status,
                        package_path,
                        notes,
                    ),
                )
                count += 1
            conn.commit()
        return count

    def update_status(
        self,
        position_id: str,
        status: str,
        *,
        package_path: str | None = None,
        notes: str | None = None,
        priority_rank: int | None = None,
    ) -> None:
        fields = ["status = ?"]
        values: list[Any] = [status]
        if package_path is not None:
            fields.append("package_path = ?")
            values.append(package_path)
        if notes is not None:
            fields.append("notes = ?")
            values.append(notes)
        if priority_rank is not None:
            fields.append("priority_rank = ?")
            values.append(priority_rank)
        values.append(position_id)
        with self.connect() as conn:
            conn.execute(
                f"UPDATE positions SET {', '.join(fields)} WHERE id = ?",
                values,
            )
            conn.commit()

    def get(self, position_id: str) -> dict[str, Any] | None:
        with self.connect() as conn:
            row = conn.execute("SELECT * FROM positions WHERE id = ?", (position_id,)).fetchone()
            return dict(row) if row else None

    def list_by_status(self, statuses: list[str]) -> list[dict[str, Any]]:
        placeholders = ",".join("?" for _ in statuses)
        with self.connect() as conn:
            rows = conn.execute(
                f"SELECT * FROM positions WHERE status IN ({placeholders}) ORDER BY score DESC",
                statuses,
            ).fetchall()
            return [dict(r) for r in rows]

    def list_all(self, limit: int | None = None) -> list[dict[str, Any]]:
        sql = (
            "SELECT * FROM positions ORDER BY "
            "CASE WHEN priority_rank IS NULL OR priority_rank = 0 THEN 1 ELSE 0 END, "
            "priority_rank, score DESC"
        )
        if limit:
            sql += f" LIMIT {int(limit)}"
        with self.connect() as conn:
            return [dict(r) for r in conn.execute(sql).fetchall()]

    def known_applied_urls(self) -> set[str]:
        with self.connect() as conn:
            hist = conn.execute("SELECT url FROM applied_history WHERE url IS NOT NULL AND url != ''").fetchall()
        return {r["url"].strip().lower() for r in hist if r["url"]}

    def known_history_keys(self) -> set[str]:
        with self.connect() as conn:
            rows = conn.execute("SELECT key FROM applied_history").fetchall()
        return {r["key"] for r in rows}

    def import_applied_history(self, entries: list[dict[str, str]]) -> int:
        from datetime import datetime, timezone

        now = datetime.now(timezone.utc).isoformat(timespec="seconds")
        count = 0
        with self.connect() as conn:
            for e in entries:
                key = e.get("key") or e.get("label") or e.get("url")
                if not key:
                    continue
                conn.execute(
                    """
                    INSERT INTO applied_history (key, label, url, imported_at)
                    VALUES (?, ?, ?, ?)
                    ON CONFLICT(key) DO UPDATE SET
                        label=excluded.label,
                        url=COALESCE(excluded.url, applied_history.url)
                    """,
                    (key.lower(), e.get("label", ""), e.get("url", ""), now),
                )
                count += 1
            conn.commit()
        return count

    def export_excel(self, path: Path | None = None) -> Path:
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        path = path or (OUTPUT_DIR / f"daily_{date.today().isoformat()}.xlsx")
        rows = self.list_all()
        wb = Workbook()
        ws = wb.active
        ws.title = "positions"
        headers = [
            "priority_rank",
            "score",
            "title",
            "institution",
            "city",
            "country",
            "region_tier",
            "pi_or_contact",
            "url",
            "source",
            "deadline",
            "posted_date",
            "match_reasons",
            "status",
            "application_mode",
            "package_path",
            "email",
            "notes",
            "id",
        ]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(h) for h in headers])
        saved = _save_workbook_resilient(wb, path)

        # Also CSV snapshot (same stem as the workbook we actually wrote).
        import csv

        csv_path = saved.with_suffix(".csv")
        tmp_csv = csv_path.with_name(csv_path.name + ".tmp")
        with tmp_csv.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            for r in rows:
                writer.writerow({h: r.get(h) for h in headers})
        _replace_resilient(tmp_csv, csv_path)
        return saved


def _save_workbook_resilient(wb: Workbook, path: Path) -> Path:
    """Save workbook; if the file is locked (open in Excel), write a sibling path."""
    tmp = path.with_name(path.name + ".tmp")
    wb.save(tmp)
    return _replace_resilient(tmp, path)


def _replace_resilient(src: Path, dest: Path) -> Path:
    """Move src onto dest. If dest is locked, write a free sibling and return that path."""
    try:
        src.replace(dest)
        return dest
    except PermissionError:
        pass
    try:
        dest.write_bytes(src.read_bytes())
        src.unlink(missing_ok=True)
        return dest
    except PermissionError:
        alt = _unlocked_sibling(dest)
        print(f"[export] {dest.name} is locked (close Excel if open); writing {alt.name} instead")
        src.replace(alt)
        return alt


def _unlocked_sibling(path: Path) -> Path:
    for i in range(1, 50):
        candidate = path.with_name(f"{path.stem}_{i}{path.suffix}")
        if not candidate.exists():
            return candidate
        try:
            with candidate.open("a"):
                return candidate
        except PermissionError:
            continue
    return path.with_name(f"{path.stem}_alt{path.suffix}")
