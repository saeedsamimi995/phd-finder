from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from phd_finder.config_loader import load_profile


URL_RE = re.compile(r"https?://[^\s<>\"]+")


def load_applied_entries(path: str | Path | None = None) -> list[dict[str, str]]:
    profile = load_profile()
    xlsx = Path(path or profile["applied_history_xlsx"])
    if not xlsx.exists():
        print(f"[history] applied workbook not found: {xlsx}")
        return []
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    entries: list[dict[str, str]] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if not cells:
                continue
            blob = " | ".join(cells)
            urls = URL_RE.findall(blob)
            label = cells[0][:180]
            entries.append(
                {
                    "key": blob[:240].lower(),
                    "label": label,
                    "url": urls[0] if urls else "",
                }
            )
    wb.close()
    return entries
